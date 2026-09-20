"""Mandatory external review: parser, read-only preview and guarded confirmation."""
import io
import uuid

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from dependencies import db
from engines.sql_utils import gd_find, gd_insert, gd_update_one
from test_bulk_student_import_atomic import _mk_school, _mk_principal, _headers, _create_excel_file
from src.modules.bulk_import.services.external_preview import parse_file, MAX_BYTES


def student(**changes):
    return {"first_name": "أحمد", "last_name": "سالم", "national_id": "1234567890",
            "grade": "الصف الأول الابتدائي", "class_name": "أ",
            "parent_phone": "0501234567", "parent_name": "سالم", **changes}


async def context():
    school = "sch_" + uuid.uuid4().hex[:10]
    await _mk_school(school)
    actor = await _mk_principal(school)
    return school, actor, _headers(actor["id"], actor["role"], school)


async def preview(client, headers, rows, kind="students", **data):
    response = await client.post(f"/bulk/preview/{kind}", headers=headers, data=data,
                                 files={"file": ("external.xlsx", _create_excel_file(rows))})
    assert response.status_code == 200, response.text
    return response.json()


def confirmation(draft, **changes):
    return {key: draft[key] for key in ("draft_id", "fingerprint", "preview_version")} | {
        "acknowledged": True, **changes}


@pytest.mark.parametrize("contents,name", [
    (b"x" * (MAX_BYTES + 1), "test.csv"), (b"PKfake", "test.csv"),
    (b"a,b\n=SUM(1),x", "test.csv"), (b"a,a\n1,2", "test.csv"),
    (b"not excel", "test.xlsx"), (b"not excel", "test.xls"),
    (b"a\n" + b"x\n" * 5001, "test.csv"),
])
def test_parser_rejects_unsafe_input(contents, name):
    with pytest.raises(HTTPException):
        parse_file(contents, name)


def test_parser_rejects_workbook_formula():
    book = Workbook()
    book.active.append(["full_name"])
    book.active.append(["=1+1"])
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(HTTPException):
        parse_file(out.getvalue(), "test.xlsx")


@pytest.mark.parametrize("kind", ["students", "teachers"])
@pytest.mark.asyncio
async def test_downloaded_official_template_accepts_replacement_rows(client, kind):
    school, actor, headers = await context()
    downloaded = await client.get(f"/bulk/template/{kind}", headers=headers)
    assert downloaded.status_code == 200, downloaded.text
    book = load_workbook(io.BytesIO(downloaded.content))
    assert set(book.sheetnames) == {"البيانات", "تعليمات"}
    sheet = book["البيانات"]
    columns = [cell.value for cell in sheet[1]]
    sheet.delete_rows(2, sheet.max_row)
    values = ({
        "الاسم الأول (مطلوب)": "طالب", "اسم الأب (مطلوب)": "أحمد",
        "اسم الجد (مطلوب)": "علي", "اسم العائلة (مطلوب)": "سالم",
        "رقم الهوية (مطلوب)": "1234567890", "الصف (مطلوب)": "الصف الأول الابتدائي",
        "الفصل (مطلوب)": "أ", "جوال ولي الأمر (مطلوب)": "0501234567",
    } if kind == "students" else {
        "الاسم الكامل (مطلوب)": "معلم أحمد سالم",
        "البريد الإلكتروني (مطلوب)": f"{uuid.uuid4().hex}@teacher.test",
        "رقم الجوال (مطلوب)": "0501234567",
    })
    sheet.append([values.get(column) for column in columns])
    # Selecting by name must work even when instructions are the first sheet.
    book.move_sheet("تعليمات", offset=-1)
    output = io.BytesIO()
    book.save(output)
    book.close()
    response = await client.post(f"/bulk/preview/{kind}", headers=headers,
                                 files={"file": ("official-template.xlsx", output.getvalue())})
    assert response.status_code == 200, response.text
    draft = response.json()
    assert draft["can_confirm"], draft
    assert draft["summary"]["total_rows"] == 1
    assert draft["rows"][0]["row"] == 2
    assert draft["rows"][0]["action"] == "create"
    assert not await gd_find(db.session, kind, {"school_id": school})


@pytest.mark.parametrize("unsafe", ["unknown_sheet", "third_sheet", "instructions_formula",
                                     "instructions_rows", "instructions_columns", "instructions_cell"])
def test_additional_sheets_cannot_bypass_workbook_safety(unsafe):
    book = Workbook()
    book.active.title = "البيانات"
    book.active.append(["full_name"])
    book.active.append(["اسم"])
    instructions = book.create_sheet("تعليمات")
    instructions.append(["تعليمات الاستيراد"])
    if unsafe == "unknown_sheet":
        instructions.title = "extra-data"
    elif unsafe == "third_sheet":
        book.create_sheet("extra-data")
    elif unsafe == "instructions_formula":
        instructions["A2"] = "=1+1"
    elif unsafe == "instructions_rows":
        instructions.cell(row=5002, column=1, value="overflow")
    elif unsafe == "instructions_columns":
        instructions.cell(row=1, column=65, value="overflow")
    else:
        instructions["A2"] = "x" * 4097
    output = io.BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(HTTPException):
        parse_file(output.getvalue(), "test.xlsx")


@pytest.mark.asyncio
async def test_preview_read_only_then_confirm_consumed(client):
    school, actor, headers = await context()
    before = {table: await gd_find(db.session, table, {"school_id": school}, limit=100)
              for table in ("students", "classes", "grade_levels", "parents", "teachers")}
    parent_accounts_before = await gd_find(db.session, "users", {"tenant_id": school, "role": "parent"})
    relationships_before = await gd_find(db.session, "guardian_links", {"tenant_id": school})
    notifications_before = await gd_find(db.session, "notifications", {"school_id": school})
    draft = await preview(client, headers, [student()])
    assert draft["can_confirm"], draft
    assert draft["rows"][0]["action"] == "create"
    assert draft["summary"]["class_create"] == 1
    for table, records in before.items():
        assert await gd_find(db.session, table, {"school_id": school}, limit=100) == records
    assert await gd_find(db.session, "users", {"tenant_id": school, "role": "parent"}) == parent_accounts_before
    assert await gd_find(db.session, "guardian_links", {"tenant_id": school}) == relationships_before
    assert await gd_find(db.session, "notifications", {"school_id": school}) == notifications_before
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1, response.text
    assert response.json()["batch_id"]
    repeated = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert repeated.status_code == 409
    assert len(await gd_find(db.session, "students", {"school_id": school})) == 1


@pytest.mark.asyncio
async def test_mixed_invalid_blocks_every_row_and_legacy_bypass(client):
    school, actor, headers = await context()
    draft = await preview(client, headers, [student(), student(national_id="bad")])
    assert not draft["can_confirm"]
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 409
    assert not await gd_find(db.session, "students", {"school_id": school})
    direct = await client.post("/bulk/import/students", headers=headers,
                               files={"file": ("external.xlsx", _create_excel_file([student()]))})
    assert direct.status_code == 409


@pytest.mark.asyncio
async def test_cancel_supersede_tamper_and_cross_actor(client):
    school, actor, headers = await context()
    first = await preview(client, headers, [student()])
    second = await preview(client, headers, [student(national_id="1234567891")])
    assert (await client.post("/bulk/confirm", headers=headers, json=confirmation(first))).status_code == 409
    assert (await client.post("/bulk/confirm", headers=headers,
                              json=confirmation(second, fingerprint="bad"))).status_code == 409
    other = await _mk_principal(school)
    other_headers = _headers(other["id"], other["role"], school)
    assert (await client.post("/bulk/confirm", headers=other_headers, json=confirmation(second))).status_code == 409
    assert (await client.post("/bulk/confirm", headers=headers,
                              json=confirmation(second, rows=[]))).status_code == 422
    assert (await client.post(f"/bulk/draft/{second['draft_id']}/discard", headers=headers)).status_code == 200
    assert (await client.post("/bulk/confirm", headers=headers, json=confirmation(second))).status_code == 409
    assert not await gd_find(db.session, "students", {"school_id": school})


@pytest.mark.asyncio
async def test_stale_plan_and_expiry(client):
    school, actor, headers = await context()
    draft = await preview(client, headers, [student()])
    await gd_insert(db.session, "students", {
        "id": str(uuid.uuid4()), "school_id": school, "full_name": "طالب موجود",
        "national_id": "1234567890", "is_active": True,
    })
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 409, response.text
    assert not await gd_find(db.session, "parents", {"school_id": school})
    await gd_update_one(db.session, "external_import_drafts", {"id": draft["draft_id"]},
                        {"$set": {"expires_at": "2000-01-01T00:00:00+00:00"}})
    assert (await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))).status_code == 409


@pytest.mark.asyncio
async def test_teacher_existing_is_conflict_never_update(client):
    school, actor, headers = await context()
    row = {"full_name": "معلم", "email": "teacher@example.test", "phone": "0501234567"}
    draft = await preview(client, headers, [row], "teachers")
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1
    conflict = await preview(client, headers, [row], "teachers")
    assert not conflict["can_confirm"]
    assert conflict["rows"][0]["action"] == "conflict"


@pytest.mark.asyncio
async def test_global_parent_reused_and_in_file_identity_conflict_blocks(client):
    school, actor, headers = await context()
    parent_id = str(uuid.uuid4())
    await gd_insert(db.session, "users", {
        "id": parent_id, "email": f"{parent_id}@parent.test", "password_hash": "test",
        "role": "parent", "full_name": "ولي أمر", "is_active": True,
    })
    draft = await preview(client, headers, [student(parent_email=f"{parent_id}@parent.test")])
    assert draft["summary"]["parent_create"] == 1
    assert draft["summary"]["parent_account_reuse"] == 1
    assert draft["rows"][0]["relationships"]["parent_account"]["key"] == parent_id
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1
    assert response.json()["created_parent_user_ids"] == []

    other_school, other_actor, other_headers = await context()
    conflict = await preview(client, other_headers, [
        student(parent_phone="0501111111", parent_email="first@guardian.test"),
        student(national_id="1234567891", parent_phone="0502222222", parent_email="second@guardian.test"),
        student(national_id="1234567892", parent_phone="0501111111", parent_email="second@guardian.test"),
    ])
    assert not conflict["can_confirm"]
    assert conflict["rows"][2]["errors"]
    assert not await gd_find(db.session, "parents", {"school_id": other_school})


@pytest.mark.asyncio
async def test_runtime_relationship_failure_rolls_back_whole_row(client, monkeypatch):
    from src.modules.bulk_import.services import student_import_service
    school, actor, headers = await context()
    draft = await preview(client, headers, [student(parent_email="row-rollback@example.test")])

    async def fail_after_relationship_writes(*args, **kwargs):
        raise RuntimeError("injected relationship persistence failure")

    monkeypatch.setattr(student_import_service, "_validate_persisted_relationships", fail_after_relationship_writes)
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0
    assert response.json()["failed"] == 1
    for table in ("students", "classes", "grade_levels", "parents"):
        assert not await gd_find(db.session, table, {"school_id": school})
    assert not await gd_find(db.session, "users", {"email": "row-rollback@example.test"})
    assert not await gd_find(db.session, "guardian_links", {"tenant_id": school})


@pytest.mark.asyncio
async def test_restore_preview_and_shared_resource_counts(client):
    school, actor, headers = await context()
    await gd_insert(db.session, "students", {
        "id": str(uuid.uuid4()), "school_id": school, "full_name": "طالب سابق",
        "national_id": "1234567890", "is_active": False,
    })
    draft = await preview(client, headers, [student(), student(national_id="1234567891")])
    assert draft["rows"][0]["action"] == "restore"
    assert draft["summary"]["restore"] == 1
    assert draft["summary"]["create"] == 1
    assert draft["summary"]["parent_create"] == 1
    assert draft["summary"]["class_create"] == 1
    assert draft["summary"]["grade_create"] == 1
    response = await client.post("/bulk/confirm", headers=headers, json=confirmation(draft))
    assert response.status_code == 200, response.text
    assert response.json()["restored"] == 1
    assert response.json()["created"] == 1
    assert response.json()["parents_created"] == 1


@pytest.mark.asyncio
async def test_database_concurrent_draft_consumption_is_once():
    """Two separate PostgreSQL connections, not a shared mocked session/lock."""
    import asyncio
    from datetime import datetime, timedelta, timezone
    from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
    from sqlalchemy.pool import NullPool
    from db import _get_async_url
    from engines.sql_utils import gd_delete_one
    from src.modules.bulk_import.services.external_preview import (
        DRAFT_COLLECTION, SOURCE, load_owned, save_payload, scope_lock,
    )
    engine = create_async_engine(_get_async_url(), poolclass=NullPool)
    draft_id = str(uuid.uuid4())
    school = "concurrency-test-" + uuid.uuid4().hex
    payload = {"id": draft_id, "source": SOURCE, "state": "live", "actor_id": "test-actor",
               "school_id": school, "expires_at": (datetime.now(timezone.utc) + timedelta(minutes=1)).isoformat()}
    try:
        async with AsyncSession(engine) as session:
            await gd_insert(session, DRAFT_COLLECTION, payload)
            await session.commit()
        ready = asyncio.Event()

        async def consume(first):
            async with AsyncSession(engine) as session:
                await scope_lock(session, school)
                try:
                    held = await load_owned(session, draft_id, "test-actor", school)
                except HTTPException as exc:
                    await session.rollback()
                    return exc.status_code
                if first:
                    ready.set()
                    await asyncio.sleep(0.1)
                held["state"] = "consumed"
                await save_payload(session, draft_id, held)
                await session.commit()
                return 200

        first = asyncio.create_task(consume(True))
        await asyncio.wait_for(ready.wait(), 5)
        second = asyncio.create_task(consume(False))
        assert sorted(await asyncio.gather(first, second)) == [200, 409]
    finally:
        async with AsyncSession(engine) as session:
            await gd_delete_one(session, DRAFT_COLLECTION, {"id": draft_id})
            await session.commit()
        await engine.dispose()