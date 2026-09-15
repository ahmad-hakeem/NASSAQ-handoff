"""Regression tests for spreadsheet formula-injection hardening."""

import csv
import io
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient

from dependencies import UserRole
from src.modules.bulk_import.controllers import bulk_import_export_routes as routes
from src.modules.bulk_import.services.student_import_service import (
    _canonical_grade,
    _normalise_columns,
    _section_from_input,
)


def _test_router(monkeypatch, frame):
    """Build the bulk router with auth/database dependencies kept local."""
    current_user = {
        "id": "serialization-test-user",
        "role": UserRole.SCHOOL_ADMIN.value,
        "tenant_id": "serialization-test-school",
        "full_name": "Serialization Test",
    }

    async def fake_export(*_args, **_kwargs):
        return frame, "export.xlsx"

    async def fake_audit_insert(*_args, **_kwargs):
        return None

    def fake_require_roles(_roles):
        async def dependency():
            return current_user

        return dependency

    monkeypatch.setattr(routes, "_export_students", fake_export)
    monkeypatch.setattr(routes, "gd_insert", fake_audit_insert)

    app = FastAPI()
    app.include_router(
        routes.setup_bulk_routes(
            db=SimpleNamespace(session=None),
            get_current_user=lambda: current_user,
            require_roles=fake_require_roles,
            UserRole=UserRole,
        )
    )
    return app


@pytest.mark.asyncio
async def test_bulk_export_neutralizes_csv_cells_and_headers(monkeypatch):
    malicious = [
        "=HYPERLINK('https://evil.invalid','x')",
        "+cmd|'/C calc'!A0",
        "-1+1",
        "@SUM(1,1)",
        "\t=evil",
        "\r=evil",
    ]
    frame = pd.DataFrame({"=header": malicious, "native": list(range(6))})
    app = _test_router(monkeypatch, frame)

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/bulk/export/students",
            params={"format": "csv"},
        )

    assert response.status_code == 200, response.text
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == ["'=header", "native"]
    assert [row[0] for row in rows[1:]] == ["'" + value for value in malicious]
    assert [row[1] for row in rows[1:]] == [str(value) for value in range(6)]


@pytest.mark.asyncio
async def test_bulk_export_neutralizes_xlsx_cells_and_preserves_numeric_cells(
    monkeypatch,
):
    malicious = [
        "=1+1",
        "+1+1",
        "-1+1",
        "@SUM(1,1)",
        "\t=evil",
        "\r=evil",
    ]
    frame = pd.DataFrame({"=header": malicious, "native": list(range(6))})
    app = _test_router(monkeypatch, frame)

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/bulk/export/students",
            params={"format": "xlsx"},
        )

    assert response.status_code == 200, response.text
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=False)
    sheet = workbook["البيانات"]
    assert sheet.cell(1, 1).value == "'=header"
    for row_number, value in enumerate(malicious, start=2):
        cell = sheet.cell(row_number, 1)
        assert cell.data_type == "s"
        # xlsxwriter stores carriage returns using Excel's XML escape
        # sequence; it remains a literal string, not a formula.
        expected = ("'" + value).replace("\r", "_x000D_")
        assert cell.value == expected
        native = sheet.cell(row_number, 2)
        assert native.value == row_number - 2
        assert native.data_type == "n"


@pytest.mark.asyncio
async def test_bulk_export_xlsx_roundtrips_canonical_grade_and_section_headers(
    monkeypatch,
):
    frame = pd.DataFrame({
        "الصف (مطلوب)": ["الصف الأول الابتدائي"],
        "الفصل (مطلوب)": ["أ"],
        "الاسم الأول": ["أحمد"],
        "اسم العائلة": ["السعيد"],
    })
    app = _test_router(monkeypatch, frame)

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/bulk/export/students",
            params={"format": "xlsx"},
        )

    assert response.status_code == 200, response.text
    imported_frame = pd.read_excel(io.BytesIO(response.content), dtype=str)
    normalised = _normalise_columns(imported_frame)
    assert normalised["grade"].tolist() == ["الصف الأول الابتدائي"]
    assert normalised["class_name"].tolist() == ["أ"]
    grade = _canonical_grade(normalised["grade"].iloc[0])
    assert grade["grade"] == 1
    assert _section_from_input(normalised["class_name"].iloc[0], grade) == "أ"