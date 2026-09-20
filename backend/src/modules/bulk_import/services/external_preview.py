"""Read-only external spreadsheet planning and transaction-bound draft lifecycle."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import secrets
import zipfile
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import TYPE_CHECKING, Literal, Optional

if TYPE_CHECKING:
    import pandas as pd
from fastapi import Depends, File, Form, Header, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, StrictBool, StrictInt, StrictStr
from sqlalchemy import select, text

from engines.sql_utils import gd_find, gd_find_one, gd_insert, gd_update_one

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 64
MAX_CELL = 4096
SOURCE = "external_template_v1"
DRAFT_COLLECTION = "external_import_drafts"


def reject(code, message, status=409):
    raise HTTPException(status_code=status, detail={"code": code, "message": message})


def _data_sheet_index(names):
    """Accept a single data sheet or the exact two-sheet official template."""
    if len(names) == 1 and names[0] != "تعليمات":
        return 0
    if len(names) == 2 and set(names) == {"البيانات", "تعليمات"}:
        return names.index("البيانات")
    raise ValueError("استخدم ورقة بيانات واحدة، أو ورقتي القالب الرسمي: البيانات وتعليمات")


def parse_file(contents: bytes, filename: str) -> pd.DataFrame:
    """Bound both compressed input and expanded workbook; never evaluate formulas."""
    import pandas as pd

    if not contents or len(contents) > MAX_BYTES:
        reject("invalid_size", "الملف فارغ أو يتجاوز 10 ميغابايت", 400)
    suffix = Path(filename or "").suffix.lower()
    rows = []
    try:
        if suffix == ".csv":
            if b"\0" in contents or contents.startswith((b"PK", b"\xd0\xcf")):
                raise ValueError("نوع الملف لا يطابق امتداده")
            reader = csv.reader(io.StringIO(contents.decode("utf-8-sig")), strict=True)
            for row in reader:
                rows.append(row)
                _check_dimensions(rows, row)
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(contents)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
                    raise ValueError("حجم المصنف بعد فك الضغط يتجاوز الحد")
                if any(e.flag_bits & 1 or "vbaproject" in e.filename.lower() or
                       "externallinks/" in e.filename.lower() or
                       e.file_size > 20 * 1024 * 1024 or
                       e.file_size > max(e.compress_size, 1) * 200 for e in entries):
                    raise ValueError("مصنف غير آمن أو مضغوط بشكل غير مقبول")
                if "[Content_Types].xml" not in archive.namelist():
                    raise ValueError("ملف XLSX غير صالح")
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(contents), read_only=True, data_only=False, keep_links=False)
            try:
                selected = _data_sheet_index(book.sheetnames)
                if len(book.worksheets) != len(book.sheetnames):
                    raise ValueError("نوع ورقة غير مدعوم")
                for index, sheet in enumerate(book.worksheets):
                    if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                        raise ValueError("تجاوز عدد الصفوف أو الأعمدة")
                    sheet_rows = []
                    # Instructions are excluded from import, not from safety checks.
                    for cells in sheet.iter_rows():
                        if any(c.data_type == "f" for c in cells):
                            raise ValueError("الصيغ الحسابية غير مسموحة؛ استخدم قيماً نصية")
                        row = [c.value for c in cells]
                        sheet_rows.append(row)
                        _check_dimensions(sheet_rows, row)
                    if index == selected:
                        rows = sheet_rows
            finally:
                book.close()
        elif suffix == ".xls":
            if not contents.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                raise ValueError("نوع الملف لا يطابق امتداده")
            import xlrd
            book = xlrd.open_workbook(file_contents=contents, on_demand=True, ragged_rows=True)
            try:
                selected = _data_sheet_index(book.sheet_names())
                for index in range(book.nsheets):
                    sheet = book.sheet_by_index(index)
                    if sheet.nrows > MAX_ROWS + 1 or sheet.ncols > MAX_COLUMNS:
                        raise ValueError("تجاوز عدد الصفوف أو الأعمدة")
                    sheet_rows = []
                    # xlrd only reads stored values, never executes macros/formulas.
                    for i in range(sheet.nrows):
                        row = sheet.row_values(i)
                        sheet_rows.append(row)
                        _check_dimensions(sheet_rows, row)
                    if index == selected:
                        rows = sheet_rows
            finally:
                book.release_resources()
        else:
            raise ValueError("يجب أن يكون الملف XLS أو XLSX أو CSV")
        if len(rows) < 2 or not any(any(v not in (None, "") for v in r) for r in rows[1:]):
            raise ValueError("الملف فارغ")
        headers = [str(v or "").strip() for v in rows[0]]
        if not all(headers) or len(headers) != len(set(headers)):
            raise ValueError("عناوين الأعمدة فارغة أو مكررة")
        if any(len(r) > len(headers) for r in rows[1:]):
            raise ValueError("عدد خلايا الصف لا يطابق العناوين")
        return pd.DataFrame(
            [[None if v is None or v == "" else str(v) for v in r] +
             [None] * (len(headers) - len(r)) for r in rows[1:]], columns=headers
        )
    except HTTPException:
        raise
    except Exception as exc:
        reject("invalid_file", f"تعذر قراءة الملف: {exc}", 400)


def _check_dimensions(rows, row):
    if len(rows) > MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
        raise ValueError("الحد الأقصى 5000 صف و64 عموداً")
    for value in row:
        if value is not None and len(str(value)) > MAX_CELL:
            raise ValueError("تجاوز طول الخلية المسموح")
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "@")):
            # Allow international phone numbers, not spreadsheet expressions.
            if value.strip().startswith("+") and value.strip()[1:].isdigit():
                continue
            raise ValueError("الصيغ الحسابية غير مسموحة")


async def resolve_student_plan(db, school_id, records, classes, grades, grade_ids, errors):
    from src.modules.bulk_import.services.student_import_service import (
        _class_canonical_grade, _class_matches, _class_key,
    )
    from services.parent_linking import _assert_parent_record_available, _assert_parent_user_available
    parents = await gd_find(db.session, "parents", {"school_id": school_id}, limit=100000)
    links = await gd_find(db.session, "guardian_links", {"tenant_id": school_id, "is_active": True}, limit=100000)
    planned_parents = {}
    planned_accounts = {}
    for item in records:
        existing = item.get("existing") or {}
        item["action"] = ("restore" if existing.get("is_active") is False or existing.get("deleted_at")
                          else "update") if existing else "create"
        grade = item["grade"]
        matches = [c for c in classes if
                   (_class_canonical_grade(c, grades, grade_ids) or {}).get("grade") == grade["grade"]
                   and _class_matches(c, item["raw_class_name"], item["section"], grade)]
        if len(matches) > 1:
            errors.append({"row": item["row_num"], "field": "class_name", "message": "يوجد أكثر من فصل مطابق؛ صحح الفصول أولاً"})
        grade_row = grades.get(str(grade["grade"]))
        class_key = str(matches[0]["id"]) if matches else "new:" + repr(_class_key(grade, item["section"]))
        parent_id = existing.get("parent_id") or next(
            (link.get("parent_id") for link in links if link.get("student_id") == existing.get("id")), None
        )
        current = next((p for p in parents if p.get("id") == parent_id), None)
        phone_match = next((p for p in parents if p.get("phone") == item["parent_phone"]), None)
        email = item.get("parent_email") or existing.get("parent_email")
        parent = phone_match or current or next((p for p in parents if email and p.get("email") == email), None)
        try:
            _assert_parent_record_available(current)
            _assert_parent_record_available(parent)
            account = None
            if parent and parent.get("user_id"):
                account = next((u for u in planned_accounts.values() if u["id"] == parent["user_id"]), None)
                if account is None:
                    account = await gd_find_one(db.session, "users", {"id": parent["user_id"]})
                _assert_parent_user_available(account)
            elif parent and parent.get("email"):
                account = planned_accounts.get(parent["email"]) or await gd_find_one(
                    db.session, "users", {"email": parent["email"], "role": "parent"})
                _assert_parent_user_available(account)
            email_user = (planned_accounts.get(email) or await gd_find_one(db.session, "users", {"email": email})) if email else None
            _assert_parent_user_available(email_user)
            if email_user and (email_user.get("role") != "parent" or
                               (account and account["id"] != email_user["id"]) or
                               (parent and not account and parent.get("email") != email)):
                raise ValueError("بريد ولي الأمر مرتبط بحساب آخر")
            account = account or email_user
            key = str(parent["id"]) if parent else planned_parents.get(
                item["parent_phone"], planned_parents.get(email, "new:" + item["parent_phone"]))
            planned_parents[item["parent_phone"]] = key
            if email:
                planned_parents[email] = key
            item["relationships"] = {
                "grade": {"action": "reuse" if grade_row else "create", "key": str(grade_row["id"]) if grade_row else "new:" + str(grade["grade"])},
                "class": {"action": "reuse" if matches else "create", "key": class_key},
                "parent": {"action": "reuse" if parent and not key.startswith("new:") else "create", "key": key},
                "parent_account": {"action": "reuse" if account and not str(account["id"]).startswith("new:") else "create",
                                   "key": str(account["id"]) if account else "new:user:" + key},
                "assignment": {"action": "assign", "key": class_key},
                "guardian_link": {"action": "link", "key": key},
            }
            if matches:
                item["relationships"]["class"]["restore"] = matches[0].get("is_active") is False
                item["relationships"]["class"]["repair_grade_link"] = (
                    not matches[0].get("grade_id") or str(matches[0].get("grade_id")).isdigit())
            item["_relationship_state"] = {
                "parent": parent, "account": account, "class": matches[0] if matches else None,
                "grade": grade_row,
            }
            effective_email = email if current else (parent or {}).get("email") or email
            effective_phone = item["parent_phone"] if current else (parent or {}).get("phone") or item["parent_phone"]
            item["relationships"]["parent"].update(email=effective_email, phone=effective_phone)
            projected_account = account or {
                "id": "new:user:" + key, "email": effective_email, "role": "parent", "is_active": True,
            }
            planned_accounts[projected_account["id"]] = projected_account
            if effective_email:
                planned_accounts[effective_email] = projected_account
            projected_parent = {
                **(parent or {}), "id": key, "phone": effective_phone, "email": effective_email,
                "user_id": projected_account["id"], "is_active": True,
            }
            parents = [p for p in parents if p.get("id") != key] + [projected_parent]
        except Exception as exc:
            errors.append({"row": item["row_num"], "field": "parent_email", "message": str(getattr(exc, "detail", exc))})
        # Bind existing state as well as IDs, so intervening edits require review.
        item["existing"] = existing
    return {"records": records}


async def plan(db, df, school_id, actor, kind):
    from src.modules.bulk_import.services.student_import_service import (
        import_students, _normalise_columns, _row_student_name, normalise_cell,
    )
    from src.modules.bulk_import.controllers.bulk_import_export_routes import _import_teachers
    errors, warnings = [], []
    if kind == "students":
        result = await import_students(db, df, school_id, actor, errors, warnings, plan_only=True)
    else:
        result = await _import_teachers(db, df, school_id, actor, errors, warnings, plan_only=True)
    records = result.get("records", [])
    indexed = {r["row_num"]: r for r in records}
    display = _normalise_columns(df)
    rows = []
    for idx, raw in df.iterrows():
        if raw.dropna().empty:
            continue
        number = int(idx) + 2
        item = indexed.get(number, {})
        row_errors = [e for e in errors if e.get("row") in (1, number)]
        fields = {str(k): normalise_cell(v) for k, v in display.loc[idx].items()}
        rows.append({
            **fields,
            **{k: v for k, v in item.items() if k != "existing" and not k.startswith("_")},
            "row": number, "name": item.get("full_name") or _row_student_name(display.loc[idx]),
            "action": ("conflict" if kind == "teachers" and any("مسجل مسبقاً" in e["message"] for e in row_errors)
                       else "error") if row_errors else item.get("action", "create"),
            "errors": row_errors, "warnings": [w for w in warnings if w.get("row") == number],
        })
    counts = Counter(r["action"] for r in rows)
    summary = {"total_rows": len(rows), "blocking_errors": len(errors),
               "error_rows": sum(bool(r["errors"]) for r in rows), "warnings": len(warnings), **counts}
    for resource in ("grade", "class", "parent", "parent_account"):
        for action in ("create", "reuse"):
            summary[f"{resource}_{action}"] = len({
                r["relationships"][resource]["key"] for r in rows if not r["errors"] and
                r.get("relationships", {}).get(resource, {}).get("action") == action
            })
    summary["assignments"] = summary["guardian_links"] = sum(
        bool(r.get("relationships")) and not r["errors"] for r in rows)
    return {"rows": rows, "summary": summary, "errors": errors, "warnings": warnings,
            "can_confirm": bool(rows) and not errors, "_records": records}


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


class Confirmation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    draft_id: StrictStr
    fingerprint: StrictStr
    preview_version: StrictInt
    acknowledged: StrictBool


async def scope_lock(session, school_id):
    from src.modules.bulk_import.services.student_import_service import acquire_school_import_lock
    await acquire_school_import_lock(session, school_id)


async def audit(session, actor, school, action, kind, summary):
    import uuid
    await gd_insert(session, "audit_logs", {
        "id": str(uuid.uuid4()), "performed_by": actor["id"],
        "action": action, "timestamp": datetime.now(timezone.utc),
        "details": {"school_id": school, "import_type": kind, "summary": summary, **summary},
    })


async def load_owned(session, draft_id, actor_id, school_id):
    from pg_models import GenericDocument
    result = await session.execute(select(GenericDocument).where(
        GenericDocument.id == draft_id, GenericDocument._collection == DRAFT_COLLECTION,
    ).with_for_update().execution_options(populate_existing=True))
    row = result.scalar_one_or_none()
    payload = dict(row.data) if row else None
    if (not payload or payload.get("source") != SOURCE or payload.get("state") != "live"
            or payload.get("actor_id") != actor_id or payload.get("school_id") != school_id
            or datetime.fromisoformat(payload["expires_at"]) <= datetime.now(timezone.utc)):
        reject("draft_unavailable", "المعاينة منتهية أو ملغاة؛ يرجى معاينة الملف مجدداً")
    return payload


async def save_payload(session, draft_id, payload):
    await gd_update_one(session, DRAFT_COLLECTION, {"id": draft_id}, {"$set": payload})


def register_external_routes(router, db, require_roles, UserRole, resolve_school):
    authorized = require_roles([UserRole.PLATFORM_ADMIN, UserRole.SCHOOL_PRINCIPAL, UserRole.SCHOOL_ADMIN])

    @router.post("/preview/{import_type}")
    async def preview(
        import_type: Literal["students", "teachers"], file: UploadFile = File(...),
        supersedes_draft_id: Optional[str] = Form(None), school_id: Optional[str] = None,
        x_school_context: Optional[str] = Header(None, alias="X-School-Context"),
        current_user: dict = Depends(authorized),
    ):
        import pandas as pd

        school = resolve_school(current_user, x_school_context, school_id)
        contents = await file.read(MAX_BYTES + 1)
        df = parse_file(contents, file.filename)
        await scope_lock(db.session, school)
        if supersedes_draft_id:
            previous = await load_owned(db.session, supersedes_draft_id, current_user["id"], school)
            previous["state"] = "superseded"
            await save_payload(db.session, supersedes_draft_id, previous)
        fingerprint = hashlib.sha256(contents).hexdigest()
        # A new successful preview supersedes every live external draft in this actor/school.
        from engines.sql_utils import gd_update_many
        await gd_update_many(db.session, DRAFT_COLLECTION,
                             {"actor_id": current_user["id"], "school_id": school, "state": "live"},
                             {"$set": {"state": "superseded"}})
        reviewed = await plan(db, df, school, current_user, import_type)
        draft_id = secrets.token_urlsafe(24)
        now = datetime.now(timezone.utc)
        expires = now + timedelta(minutes=30)
        public = {k: v for k, v in reviewed.items() if not k.startswith("_")}
        public.update(draft_id=draft_id, fingerprint=fingerprint, preview_version=1,
                      expires_at=expires.isoformat(), import_type=import_type)
        payload = {"source": SOURCE, "state": "live", "fingerprint": fingerprint, "version": 1,
                   "actor_id": current_user["id"], "school_id": school, "expires_at": expires.isoformat(),
                   "kind": import_type, "filename": file.filename, "columns": list(df.columns),
                   "data": df.astype(object).where(pd.notna(df), None).values.tolist(), "plan_hash": digest(reviewed),
                   "preview": public}
        await gd_insert(db.session, DRAFT_COLLECTION, {"id": draft_id, **payload})
        await audit(db.session, current_user, school, "external_import_preview", import_type, reviewed["summary"])
        return public

    @router.post("/draft/{draft_id}/discard")
    async def discard(draft_id: str, school_id: Optional[str] = None,
                      x_school_context: Optional[str] = Header(None, alias="X-School-Context"),
                      current_user: dict = Depends(authorized)):
        import pandas as pd

        school = resolve_school(current_user, x_school_context, school_id)
        await scope_lock(db.session, school)
        payload = await load_owned(db.session, draft_id, current_user["id"], school)
        payload["state"] = "discarded"
        await save_payload(db.session, draft_id, payload)
        return {"discarded": True}

    @router.post("/confirm")
    async def confirm(body: Confirmation, school_id: Optional[str] = None,
                      x_school_context: Optional[str] = Header(None, alias="X-School-Context"),
                      current_user: dict = Depends(authorized)):
        school = resolve_school(current_user, x_school_context, school_id)
        if body.acknowledged is not True:
            reject("acknowledgement_required", "يجب الموافقة الصريحة على إنشاء وتحديث السجلات والعلاقات")
        await scope_lock(db.session, school)
        payload = await load_owned(db.session, body.draft_id, current_user["id"], school)
        if body.fingerprint != payload["fingerprint"] or body.preview_version != payload["version"]:
            reject("draft_mismatch", "بيانات التأكيد لا تطابق المعاينة")
        if not payload["preview"]["can_confirm"]:
            reject("blocking_errors", "صحح جميع الأخطاء ثم أعد المعاينة")
        # Ordinary CRUD does not take our school advisory lock. Prevent a writer
        # from changing a reviewed identity/relationship between replan and
        # execution, including global parent accounts. These are transaction
        # locks, not process-local mutexes.
        await db.session.execute(text(
            "LOCK TABLE students, teachers, classes, grade_levels, parents, users, generic_documents "
            "IN SHARE ROW EXCLUSIVE MODE"
        ))
        df = pd.DataFrame(payload["data"], columns=payload["columns"])
        reviewed = await plan(db, df, school, current_user, payload["kind"])
        if not reviewed["can_confirm"] or digest(reviewed) != payload["plan_hash"]:
            reject("stale_plan", "تغيرت البيانات أو العلاقات؛ أعد المعاينة قبل الاستيراد")
        from src.modules.bulk_import.controllers.bulk_import_export_routes import _import_students, _import_teachers
        errors, warnings = [], []
        if payload["kind"] == "students":
            result = await _import_students(db, df, school, current_user, errors, warnings, filename=payload["filename"])
        else:
            result = await _import_teachers(db, df, school, current_user, errors, warnings)
        result.update(total_rows=len(reviewed["rows"]), errors=errors, warnings=warnings,
                      success=not errors, message="اكتمل الاستيراد" if not errors else "اكتمل الاستيراد مع أخطاء")
        payload["state"] = "consumed"
        payload["result"] = result
        await save_payload(db.session, body.draft_id, payload)
        await audit(db.session, current_user, school, f"bulk_import_{payload['kind']}", payload["kind"],
                    {key: value for key, value in result.items()
                     if isinstance(value, (int, bool)) or key == "batch_id"})
        return result