"""Create and remove a disposable live-browser principal import fixture.

This helper owns only rows whose school id and account addresses carry the
    ``live-import-qa-`` namespace.  It deliberately does not use an API route:
the browser spec uploads the workbook through the real principal UI and the
live FastAPI import route.  The generated state/workbook are written below
``/tmp`` by default and contain synthetic ``example.com`` data only.

Usage (the Playwright spec invokes these modes):

    python backend/scripts/live_principal_import_qa_fixture.py --seed
    python backend/scripts/live_principal_import_qa_fixture.py --cleanup
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import secrets
import shutil
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_STATE = Path("/tmp/nassaq_live_principal_import_qa.json")
DEFAULT_WORKBOOK = Path("/tmp/nassaq_live_principal_import_qa.xlsx")
STATE_PATH = Path(os.environ.get("LIVE_IMPORT_QA_STATE", DEFAULT_STATE))
WORKBOOK_PATH = Path(os.environ.get("LIVE_IMPORT_QA_WORKBOOK", DEFAULT_WORKBOOK))
SCHOOL_PREFIX = "live-import-qa-"
EMAIL_PREFIX = "live-import-qa-principal-"

GRADE_LABELS = (
    "الصف الأول الابتدائي",
    "الصف الثاني الابتدائي",
    "الصف الثالث الابتدائي",
    "الصف الرابع الابتدائي",
    "الصف الخامس الابتدائي",
    "الصف السادس الابتدائي",
    "الصف الأول المتوسط",
    "الصف الثاني المتوسط",
    "الصف الثالث المتوسط",
    "الصف الأول الثانوي",
)
SECTIONS = ("أ", "ب", "ج", "د", "هـ", "و", "أ", "ب", "ج", "د")
HEADERS = (
    "الاسم الأول (مطلوب)",
    "اسم الأب (مطلوب)",
    "اسم الجد(مطلوب)",
    "اسم العائلة (مطلوب)",
    "رقم الهوية (مطلوب)",
    "تاريخ الميلاد (YYYY-MM-DD)",
    "الجنس (ذكر/أنثى)",
    "الصف(مطلوب)",
    "الفصل(مطلوب)",
    "اسم ولي الأمر(مطلوب)",
    "جوال ولي الأمر (مطلوب)",
    "بريد ولي الأمر",
    "الحالة الصحية",
    "ملاحظات",
)


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _read_state() -> dict:
    if not STATE_PATH.is_file():
        raise SystemExit("live import QA state file is missing")
    try:
        state = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except Exception as exc:  # pragma: no cover - defensive cleanup guard
        raise SystemExit("live import QA state file is not valid JSON") from exc
    school_id = str(state.get("school_id") or "")
    principal_email = str(state.get("principal_email") or "")
    if not school_id.startswith(SCHOOL_PREFIX) or not principal_email.startswith(
        EMAIL_PREFIX
    ):
        raise SystemExit("refusing to use a state file outside the QA namespace")
    return state


def _synthetic_rows(run_token: str) -> tuple[list[dict], list[str], list[str]]:
    # The shape mirrors the supplied ten-row workbook: ten grades, six section
    # labels, and ten distinct grade/section pairs.  Every row has its own
    # synthetic parent pair, matching the uploaded artifact's ten pair rows.
    numeric_token = str(int(run_token[:10], 16) % 10_000_000).zfill(7)
    rows: list[dict] = []
    student_names: list[str] = []
    class_names: list[str] = []
    for index, (grade, section) in enumerate(zip(GRADE_LABELS, SECTIONS), start=1):
        first = f"اختبار حي {run_token[:6]} {index}"
        parent = f"ولي اختبار حي {run_token[:6]} {index}"
        student_names.append(f"{first} أب جد عائلة")
        class_names.append(f"{grade} - {section}")
        # These are intentionally synthetic, non-routable values.  The
        # importer only requires a sufficiently long phone and a ten-digit
        # student identifier.
        national_id = f"9{numeric_token}{index:02d}"[-10:]
        parent_phone = f"050{numeric_token}{index:01d}"
        rows.append(
            {
                HEADERS[0]: first,
                HEADERS[1]: "أب",
                HEADERS[2]: "جد",
                HEADERS[3]: "عائلة",
                HEADERS[4]: national_id,
                HEADERS[5]: f"201{index % 10}-01-{index:02d}",
                HEADERS[6]: "ذكر" if index % 2 else "أنثى",
                HEADERS[7]: grade,
                HEADERS[8]: section,
                HEADERS[9]: parent,
                HEADERS[10]: parent_phone,
                HEADERS[11]: f"live-import-qa-{run_token[:8]}-{index}@example.com",
                HEADERS[12]: "لا توجد",
                HEADERS[13]: "بيانات اصطناعية لاختبار QA",
            }
        )
    return rows, student_names, class_names


def _write_workbook(rows: list[dict]) -> None:
    from openpyxl import load_workbook

    template = ROOT / "attached_assets/0_template_NEW_with_10_students_1789490500064.xlsx"
    if not template.is_file():
        raise SystemExit("uploaded workbook template is missing")
    # Start from the exact uploaded workbook so the instruction sheet and
    # column structure remain authentic, then replace every data cell with
    # synthetic values.  No source workbook PII is written to the fixture.
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, WORKBOOK_PATH)
    workbook = load_workbook(WORKBOOK_PATH)
    sheet = workbook["البيانات"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row_number, row in enumerate(rows, start=2):
        for column_number, header in enumerate(HEADERS, start=1):
            sheet.cell(row=row_number, column=column_number).value = row[header]
    workbook.save(WORKBOOK_PATH)
    try:
        WORKBOOK_PATH.chmod(0o600)
    except OSError:
        pass


async def _seed() -> None:
    if STATE_PATH.exists():
        raise SystemExit("refusing to seed while a prior QA state file exists")
    if WORKBOOK_PATH.exists():
        raise SystemExit("refusing to seed while a prior QA workbook exists")

    # Importing backend modules after the path setup keeps this script usable
    # from the repository root without changing application source code.
    backend_dir = ROOT / "backend"
    if str(backend_dir) not in sys.path:
        sys.path.insert(0, str(backend_dir))
    os.environ.setdefault("TESTING", "1")
    from dependencies import hash_password
    from db import async_session_factory
    from engines.sql_utils import gd_find_one, gd_insert

    run_token = secrets.token_hex(10)
    school_id = f"{SCHOOL_PREFIX}{run_token}"
    principal_id = str(uuid.uuid4())
    principal_email = f"{EMAIL_PREFIX}{run_token}@example.com"
    password = secrets.token_urlsafe(24)
    now = _now_iso()
    rows, student_names, class_names = _synthetic_rows(run_token)
    _write_workbook(rows)

    state = {
        "school_id": school_id,
        "principal_id": principal_id,
        "principal_email": principal_email,
        "principal_password": password,
        "workbook": str(WORKBOOK_PATH),
        "student_names": student_names,
        "class_names": class_names,
        "expected_rows": 10,
        "expected_grade_count": 10,
        "expected_section_count": 6,
        "expected_pair_count": 10,
    }

    # Write ownership metadata before touching the database.  If a later
    # insert/commit or the browser setup fails, cleanup still knows exactly
    # which generated tenant it may remove.
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        STATE_PATH.write_text(
            json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        STATE_PATH.chmod(0o600)
    except OSError:
        if WORKBOOK_PATH.exists():
            WORKBOOK_PATH.unlink()
        raise

    try:
        async with async_session_factory() as session:
            # Safety checks happen before any write.  Generated identifiers
            # are random, and a collision aborts rather than touching a row.
            if await gd_find_one(session, "schools", {"id": school_id}):
                raise RuntimeError("generated school id unexpectedly exists")
            if await gd_find_one(session, "users", {"email": principal_email}):
                raise RuntimeError("generated principal email unexpectedly exists")
            await gd_insert(
                session,
                "schools",
                {
                    "id": school_id,
                    "name": f"Live import QA {run_token[:8]}",
                    "name_en": f"Live import QA {run_token[:8]}",
                    "code": f"LIVE-QA-{run_token[:8]}",
                    "country": "SA",
                    "language": "ar",
                    "calendar_system": "hijri_gregorian",
                    "school_type": "public",
                    "tenant_type": "development",
                    "status": "active",
                    "principal_id": principal_id,
                    "principal_name": "مدير اختبار حي",
                    "principal_email": principal_email,
                    "current_students": 0,
                    "current_teachers": 0,
                    "created_at": now,
                    "updated_at": now,
                },
            )
            await gd_insert(
                session,
                "users",
                {
                    "id": principal_id,
                    "email": principal_email,
                    "full_name": "مدير اختبار حي",
                    "password_hash": hash_password(password),
                    "role": "school_principal",
                    "linked_roles": ["school_principal"],
                    "tenant_id": school_id,
                    "primary_tenant_id": school_id,
                    "status": "active",
                    "is_active": True,
                    "must_change_password": False,
                    "preferred_language": "ar",
                    "email_verified": True,
                    "mfa_enrolled_at": now,
                    "created_at": now,
                    "updated_at": now,
                },
            )
            await session.commit()
    except Exception:
        # The marker was written before the transaction so this cleanup also
        # covers an error after a successful commit.
        try:
            await _cleanup()
        except Exception:
            # Preserve the marker/workbook if cleanup itself is unavailable;
            # a caller can safely retry --cleanup with the same ownership
            # metadata instead of losing track of generated rows.
            pass
        raise


async def _cleanup() -> None:
    state = _read_state()
    backend_dir = ROOT / "backend"
    if str(backend_dir) not in sys.path:
        sys.path.insert(0, str(backend_dir))
    os.environ.setdefault("TESTING", "1")
    from db import async_session_factory
    from sqlalchemy import text

    school_id = state["school_id"]
    principal_id = state["principal_id"]
    principal_email = state["principal_email"]
    async with async_session_factory() as session:
        try:
            # Remove only rows belonging to this generated tenant/account.
            # The table/column checks make cleanup safe across development
            # schemas that predate one of the optional audit/session tables.
            table_rows = (
                await session.execute(
                    text(
                        """
                        SELECT table_name
                        FROM information_schema.tables
                        WHERE table_schema = current_schema()
                        """
                    )
                )
            ).scalars().all()
            tables = set(map(str, table_rows))

            async def columns(table: str) -> set[str]:
                result = await session.execute(
                    text(
                        """
                        SELECT column_name
                        FROM information_schema.columns
                        WHERE table_schema = current_schema()
                          AND table_name = :table
                        """
                    ),
                    {"table": table},
                )
                return {str(value) for value in result.scalars().all()}

            user_ids_result = await session.execute(
                text(
                    """
                    SELECT id FROM users
                    WHERE id = :principal_id
                       OR tenant_id = :school_id
                       OR email = :principal_email
                    """
                ),
                {
                    "principal_id": principal_id,
                    "school_id": school_id,
                    "principal_email": principal_email,
                },
            )
            user_ids = {str(value) for value in user_ids_result.scalars().all()}
            user_ids.add(principal_id)

            # Login creates session/audit records before the import.  Remove
            # those generated records explicitly rather than relying only on
            # school CASCADE behavior.
            for table in (
                "user_sessions",
                "revoked_tokens",
                "mfa_factors",
                "mfa_recovery_codes",
                "security_events",
            ):
                if table not in tables:
                    continue
                table_columns = await columns(table)
                if "user_id" in table_columns and user_ids:
                    for user_id in user_ids:
                        await session.execute(
                            text(f"DELETE FROM {table} WHERE user_id = :user_id"),
                            {"user_id": user_id},
                        )

            if "audit_logs" in tables:
                audit_columns = await columns("audit_logs")
                clauses: list[str] = []
                params: dict = {"school_id": school_id}
                if "school_id" in audit_columns:
                    clauses.append("school_id = :school_id")
                if "tenant_id" in audit_columns:
                    clauses.append("tenant_id = :school_id")
                if "performed_by" in audit_columns:
                    for index, user_id in enumerate(user_ids):
                        key = f"audit_performed_by_{index}"
                        clauses.append(f"performed_by = :{key}")
                        params[key] = user_id
                if "user_id" in audit_columns:
                    for index, user_id in enumerate(user_ids):
                        key = f"audit_user_{index}"
                        clauses.append(f"user_id = :{key}")
                        params[key] = user_id
                if "details" in audit_columns:
                    clauses.append("details ->> 'school_id' = :school_id")
                if clauses:
                    await session.execute(
                        text(f"DELETE FROM audit_logs WHERE {' OR '.join(clauses)}"),
                        params,
                    )

            # guardian_links are intentionally a GenericDocument collection
            # (there is no physical guardian_links table in the mapped schema).
            # Remove these links before users/students so the generated parent
            # references cannot survive tenant cleanup.
            if "generic_documents" in tables:
                generic_columns = await columns("generic_documents")
                if {"collection", "data"}.issubset(generic_columns):
                    await session.execute(
                        text(
                            """
                            DELETE FROM generic_documents
                            WHERE collection = 'guardian_links'
                              AND (
                                data ->> 'tenant_id' = :school_id
                                OR data ->> 'school_id' = :school_id
                              )
                            """
                        ),
                        {"school_id": school_id},
                    )

            # Explicit rows named in the QA cleanup contract.
            for table, column in (
                ("guardian_links", "tenant_id"),
                ("bulk_import_batches", "school_id"),
                ("students", "school_id"),
                ("classes", "school_id"),
                ("grade_levels", "school_id"),
                ("parents", "school_id"),
            ):
                if table in tables and column in await columns(table):
                    await session.execute(
                        text(f"DELETE FROM {table} WHERE {column} = :school_id"),
                        {"school_id": school_id},
                    )

            if "users" in tables:
                user_columns = await columns("users")
                predicates = ["id = :principal_id", "email = :principal_email"]
                params = {
                    "principal_id": principal_id,
                    "principal_email": principal_email,
                }
                if "tenant_id" in user_columns:
                    predicates.append("tenant_id = :school_id")
                    params["school_id"] = school_id
                await session.execute(
                    text(f"DELETE FROM users WHERE {' OR '.join(predicates)}"),
                    params,
                )
            if "schools" in tables:
                await session.execute(
                    text("DELETE FROM schools WHERE id = :school_id"),
                    {"school_id": school_id},
                )
            await session.commit()
        except Exception:
            await session.rollback()
            raise

    # Verify the exact generated tenant/account no longer exists before
    # removing the state marker.  A failed verification keeps the marker so a
    # caller can retry cleanup safely rather than losing ownership metadata.
    async with async_session_factory() as verify_session:
        school_count = (
            await verify_session.execute(
                text("SELECT count(*) FROM schools WHERE id = :school_id"),
                {"school_id": school_id},
            )
        ).scalar_one()
        user_count = (
            await verify_session.execute(
                text(
                    """
                    SELECT count(*) FROM users
                    WHERE id = :principal_id
                       OR email = :principal_email
                       OR tenant_id = :school_id
                    """
                ),
                {
                    "principal_id": principal_id,
                    "principal_email": principal_email,
                    "school_id": school_id,
                },
            )
        ).scalar_one()
        if school_count or user_count:
            raise SystemExit("live import QA cleanup verification failed")
        generic_tables = (
            await verify_session.execute(
                text(
                    """
                    SELECT table_name
                    FROM information_schema.tables
                    WHERE table_schema = current_schema()
                    """
                )
            )
        ).scalars().all()
        if "generic_documents" in set(generic_tables):
            generic_count = (
                await verify_session.execute(
                    text(
                        """
                        SELECT count(*) FROM generic_documents
                        WHERE collection = 'guardian_links'
                          AND (
                            data ->> 'tenant_id' = :school_id
                            OR data ->> 'school_id' = :school_id
                          )
                        """
                    ),
                    {"school_id": school_id},
                )
            ).scalar_one()
            if generic_count:
                raise SystemExit("live import QA generic guardian cleanup failed")

    if WORKBOOK_PATH.exists():
        WORKBOOK_PATH.unlink()
    STATE_PATH.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--seed", action="store_true")
    group.add_argument("--cleanup", action="store_true")
    args = parser.parse_args()
    if args.seed:
        asyncio.run(_seed())
    else:
        asyncio.run(_cleanup())


if __name__ == "__main__":
    main()